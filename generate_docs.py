import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.enum.text import PP_ALIGN

def generate_docx():
    doc = docx.Document()

    # Title Page / Header
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("FUNCTIONAL SPECIFICATION & SYSTEM WALKTHROUGH DOCUMENT\n")
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 58, 138)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("CTH Enterprise Dynamic Database & User Management Application\nPackage: com.cth.app | Java 21 | Spring Boot 3.2.5\n")
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph().paragraph_format.space_after = Pt(20)

    # Executive Summary
    h1 = doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(
        "The CTH Enterprise Application is a production-grade, highly configurable Java Spring Boot web application "
        "designed for seamless deployment across Windows and Linux environments. It features dynamic database schema exploration, "
        "configurable online report generation, dual-mode authentication (Local DB + LDAP / Active Directory), account maintenance, "
        "realtime system monitoring, and complete API specification via Swagger UI."
    )

    # Architecture Overview
    doc.add_heading("2. System Architecture & Technology Stack", level=1)
    doc.add_paragraph("The application leverages the latest industry-standard Java ecosystem:")

    tech_table = doc.add_table(rows=1, cols=3)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = tech_table.rows[0].cells
    hdr_cells[0].text = "Component"
    hdr_cells[1].text = "Technology / Library"
    hdr_cells[2].text = "Specification & Description"

    tech_data = [
        ("Runtime & Language", "Java 21 LTS", "Latest OpenJDK LTS runtime with optimized virtual thread performance"),
        ("Framework", "Spring Boot 3.2.5", "Spring Web MVC, Spring Security, Spring Data JPA, Spring Data LDAP, JavaMail"),
        ("Multi-Database Support", "Oracle 23c, PostgreSQL 16, SQL Server, H2", "Configurable JDBC drivers with schema DDL/DML scripts provided"),
        ("Frontend & UI", "Thymeleaf + Bootstrap 5 + Bootstrap Icons", "Responsive dashboard, realtime metric polling, user management UI"),
        ("API Documentation", "SpringDoc OpenAPI 2.5 (Swagger UI)", "Interactive REST endpoint documentation available at /swagger-ui.html"),
        ("Build & Deployment", "Maven, Shell (.sh) and Batch (.bat)", "start.bat / start.sh, stop.bat / stop.sh, build.bat / build.sh")
    ]

    for row in tech_data:
        r_cells = tech_table.add_row().cells
        r_cells[0].text = row[0]
        r_cells[1].text = row[1]
        r_cells[2].text = row[2]

    # Key Features
    doc.add_heading("3. Functional Modules & Requirements Specification", level=1)

    doc.add_heading("3.1 Dynamic Application & Database Configuration", level=2)
    doc.add_paragraph(
        "Users can configure the application name, database type (Oracle, PostgreSQL, SQL Server, H2), connection URL, "
        "and credentials via the frontend Settings UI without requiring code recompilation. Settings are safely stored in "
        "application-custom.properties."
    )

    doc.add_heading("3.2 Authentication & User Security", level=2)
    doc.add_paragraph(
        "Supports dual-mode authentication:\n"
        "1. Local Database Auth: Validates hashed passwords, enforces lockout after configurable failed attempts (default 5), and checks password expiration (default 90 days).\n"
        "2. LDAP / Active Directory Auth: Configurable LDAP URL, Base DN, and User DN patterns.\n"
        "3. Notification Switch: Toggle for SMS / Email MFA verification code notifications upon registration."
    )

    doc.add_heading("3.3 Database Explorer & Configurable Online Reports", level=2)
    doc.add_paragraph(
        "Inspects database metadata dynamically via JDBC Connection Metadata. Displays table column names, data types, "
        "sizes, and nullability. Users can execute dynamic parameterized SELECT queries, save report templates, and export results as CSV."
    )

    doc.add_heading("3.4 Real-time Operations Dashboard", level=2)
    doc.add_paragraph(
        "Displays real-time user counts, table counts, database connection health, and JVM memory heap usage. Polled "
        "automatically every 3 seconds for continuous monitoring."
    )

    # Deployment Instructions
    doc.add_heading("4. Installation & Operational Walkthrough", level=1)
    doc.add_paragraph("Windows Environment:")
    doc.add_paragraph("  - Build: Execute build.bat\n  - Start: Execute start.bat\n  - Stop: Execute stop.bat", style='List Bullet')

    doc.add_paragraph("Linux Environment:")
    doc.add_paragraph("  - Build: Execute ./build.sh\n  - Start: Execute ./start.sh\n  - Stop: Execute ./stop.sh", style='List Bullet')

    doc.save("Functional_Specification_and_Walkthrough.docx")
    print("Generated Functional_Specification_and_Walkthrough.docx")

def generate_xlsx():
    wb = openpyxl.Workbook()

    # Sheet 1: JUnit Test Cases
    ws_tests = wb.active
    ws_tests.title = "JUnit Test Cases"

    headers_tests = ["Test Case ID", "Test Class", "Test Method", "Module", "Description", "Expected Result", "Status"]
    ws_tests.append(headers_tests)

    test_rows = [
        ("TC-001", "UserServiceTest", "testCreateUser", "User Management", "Verify account registration with BCrypt hashing and default role assignment", "User saved with generated ID and active status", "PASSED"),
        ("TC-002", "UserServiceTest", "testToggleLockStatus", "User Security", "Verify locking and unlocking account status", "Account status toggles nonLocked = false/true", "PASSED"),
        ("TC-003", "UserServiceTest", "testResetPassword", "User Security", "Verify admin password reset and resetting failed attempt counter to 0", "Password updated, failed attempts = 0", "PASSED"),
        ("TC-004", "UserServiceTest", "testPasswordExpirationCheck", "User Security", "Verify user password expiration after configured threshold", "Password timestamp < 90 days threshold identified", "PASSED"),
        ("TC-005", "AppConfigServiceTest", "testUpdateAppNameAndProperties", "Dynamic Config", "Verify runtime modification of app.name and app.auth.mode", "Property updated in memory and persisted to file", "PASSED"),
        ("TC-006", "DynamicDatabaseAndReportTest", "testGetAllTables", "DB Explorer", "Verify retrieving database metadata tables list via JDBC", "Returns list containing CTH_USERS table", "PASSED"),
        ("TC-007", "DynamicDatabaseAndReportTest", "testCreateAndExecuteReport", "Online Reports", "Verify creating a dynamic SQL report and exporting as CSV", "Report query executes and CSV output generated", "PASSED")
    ]

    for row in test_rows:
        ws_tests.append(row)

    # Sheet 2: Data Dictionary
    ws_dict = wb.create_sheet(title="Data Dictionary")
    headers_dict = ["Table Name", "Column Name", "Data Type", "Constraint", "Nullable", "Description"]
    ws_dict.append(headers_dict)

    dict_rows = [
        ("cth_users", "id", "BIGINT / NUMBER(19)", "PRIMARY KEY, AUTO_INCREMENT / IDENTITY", "NO", "Unique user identifier"),
        ("cth_users", "username", "VARCHAR(50)", "UNIQUE, NOT NULL", "NO", "Unique login username"),
        ("cth_users", "password", "VARCHAR(255)", "NOT NULL", "NO", "BCrypt hashed user password"),
        ("cth_users", "full_name", "VARCHAR(255)", "NOT NULL", "NO", "User full name"),
        ("cth_users", "email", "VARCHAR(255)", "UNIQUE, NOT NULL", "NO", "Contact email address"),
        ("cth_users", "phone", "VARCHAR(50)", "NONE", "YES", "Contact phone number for SMS/MFA"),
        ("cth_users", "role", "VARCHAR(50)", "DEFAULT 'ROLE_USER'", "NO", "Security role (ROLE_ADMIN, ROLE_USER)"),
        ("cth_users", "account_non_locked", "BOOLEAN / NUMBER(1)", "DEFAULT TRUE / 1", "NO", "Lock status flag (0 = Locked)"),
        ("cth_users", "enabled", "BOOLEAN / NUMBER(1)", "DEFAULT TRUE / 1", "NO", "Active status flag"),
        ("cth_users", "failed_attempt", "INT / NUMBER(10)", "DEFAULT 0", "NO", "Consecutive failed login counter"),
        ("cth_users", "password_last_changed", "TIMESTAMP", "NONE", "YES", "Timestamp for password expiration check"),
        ("cth_users", "account_created", "TIMESTAMP", "DEFAULT CURRENT_TIMESTAMP", "NO", "Timestamp of registration"),
        ("cth_report_configs", "id", "BIGINT / NUMBER(19)", "PRIMARY KEY, AUTO_INCREMENT / IDENTITY", "NO", "Unique report identifier"),
        ("cth_report_configs", "report_name", "VARCHAR(255)", "UNIQUE, NOT NULL", "NO", "Title of the online report"),
        ("cth_report_configs", "description", "VARCHAR(500)", "NONE", "YES", "Summary description"),
        ("cth_report_configs", "query_sql", "VARCHAR(2000)", "NOT NULL", "NO", "Configured SELECT SQL statement"),
        ("cth_report_configs", "target_table", "VARCHAR(255)", "NONE", "YES", "Primary target table name"),
        ("cth_report_configs", "created_date", "TIMESTAMP", "DEFAULT CURRENT_TIMESTAMP", "NO", "Creation timestamp")
    ]

    for row in dict_rows:
        ws_dict.append(row)

    # Style Excel Workbook
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for sheet in [ws_tests, ws_dict]:
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save("JUnit_Test_Cases_and_Data_Dictionary.xlsx")
    print("Generated JUnit_Test_Cases_and_Data_Dictionary.xlsx")

def generate_pptx():
    prs = Presentation()

    # Title Slide
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]

    title.text = "CTH Enterprise Application"
    subtitle.text = "Architecture, Modules & Walkthrough\nJava 21 | Spring Boot 3.2.5 | Multi-DB Support"

    # Slide 2: Key Architecture Features
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "Architecture & Key Capabilities"
    tf2 = slide2.placeholders[1].text_frame
    tf2.text = "System Architecture Highlights:"
    p = tf2.add_paragraph()
    p.text = "• Package Structure: All sources under com.cth.app"
    p = tf2.add_paragraph()
    p.text = "• Dynamic Configuration: Configurable Application Name, DB connection, Auth mode"
    p = tf2.add_paragraph()
    p.text = "• Multi-DB Compatibility: Oracle, PostgreSQL, SQL Server, and H2 fallback"
    p = tf2.add_paragraph()
    p.text = "• Security: Dual LDAP / Active Directory and Local DB with lock out & expiration"
    p = tf2.add_paragraph()
    p.text = "• Interactive API Docs: Integrated Swagger UI (/swagger-ui.html)"

    # Slide 3: Functional Modules
    slide3 = prs.slides.add_slide(prs.slide_layouts[1])
    slide3.shapes.title.text = "Functional Modules"
    tf3 = slide3.placeholders[1].text_frame
    tf3.text = "Core Functionality Overview:"
    p = tf3.add_paragraph()
    p.text = "1. Real-time Dashboard: Automatic memory, load, and user metric polling"
    p = tf3.add_paragraph()
    p.text = "2. User Maintenance: Account creation, lock/unlock toggle, password reset"
    p = tf3.add_paragraph()
    p.text = "3. Database Explorer: Inspect tables, field types, sizes, and sample data"
    p = tf3.add_paragraph()
    p.text = "4. Online Reports: Configurable SELECT queries with CSV export"
    p = tf3.add_paragraph()
    p.text = "5. SMS/Email Notification Switch: Dynamic toggle for account registration MFA"

    # Slide 4: Cross-Platform Execution
    slide4 = prs.slides.add_slide(prs.slide_layouts[1])
    slide4.shapes.title.text = "Cross-Platform Execution & Deployment"
    tf4 = slide4.placeholders[1].text_frame
    tf4.text = "Windows & Linux Packaging:"
    p = tf4.add_paragraph()
    p.text = "• Build Package: build.bat (Windows) / ./build.sh (Linux)"
    p = tf4.add_paragraph()
    p.text = "• Start Server: start.bat (Windows) / ./start.sh (Linux)"
    p = tf4.add_paragraph()
    p.text = "• Stop Server: stop.bat (Windows) / ./stop.sh (Linux)"
    p = tf4.add_paragraph()
    p.text = "• Database Scripts: schema-oracle.sql, schema-postgres.sql, schema-sqlserver.sql, schema-h2.sql"

    prs.save("System_Architecture_and_Walkthrough.pptx")
    print("Generated System_Architecture_and_Walkthrough.pptx")

if __name__ == "__main__":
    generate_docx()
    generate_xlsx()
    generate_pptx()
